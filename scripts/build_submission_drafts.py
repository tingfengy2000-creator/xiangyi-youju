"""Fill the user's three official DOCX templates without rewriting unrelated parts.

Run with the bundled document Python runtime. The original Downloads files are
read-only inputs; the three outputs remain drafts until visual QA and missing
evidence have been completed. No network, system installation or model calls.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "submission"
EXPECTED_HASHES = {
    0: "56e3a668957664bf11a0790055b8003fa1dcf7e04a110ae8902558c8cfed0766",
    1: "c1f280161f0037b477528b65372672c003806457f4923be773ef456cd1010697",
    2: "bf0347f12f93bba3280e2c0ca384727fd7ad2360c64201fb6d1c1cb648f655cf",
    3: "777f66a29204f6a9470da43f0b23d1e3c76f08a4891c08718c81c27dfd91e01b",
}
LIMITS = {
    "form": [("项目背景", ["background"], 300), ("立项思路", ["idea"], 300),
             ("解决方案", ["solution"], 600), ("商业模式和预期效益", ["benefits"], 300)],
    "description": [("立项依据", ["basis"], 2000),
                    ("项目创新内容", ["overall", "feasibility", "innovation"], 3000),
                    ("实施方案", ["implementation"], 3000),
                    ("应用前景分析", ["prospects"], 500)],
    "business": [("项目方案概述", ["overview"], 200), ("项目团队", ["team"], 200),
                 ("项目产品（服务）化", ["features", "product_plan"], 2000),
                 ("项目产品（服务）市场与竞争", ["market", "competition", "risks"], 2000),
                 ("商业模式", ["development", "marketing", "revenue", "venture"], 2000),
                 ("预期经济效益分析", ["economics"], 500)],
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_paragraph(paragraph, text: str) -> None:
    """Replace a slot while retaining the paragraph and its first run style."""
    first_run = paragraph.find(qn("w:r"))
    run_properties = copy.deepcopy(first_run.find(qn("w:rPr"))) if first_run is not None else None
    for child in list(paragraph):
        if child.tag not in {qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")}:
            paragraph.remove(child)
    run = OxmlElement("w:r")
    if run_properties is not None:
        run.append(run_properties)
    else:
        run_properties = OxmlElement("w:rPr")
        run.append(run_properties)
    # Completed content should not inherit the grey instructional placeholder ink.
    for color in list(run_properties.findall(qn("w:color"))):
        run_properties.remove(color)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    run_properties.append(color)
    value = OxmlElement("w:t")
    value.set(qn("xml:space"), "preserve")
    value.text = text
    run.append(value)
    paragraph.append(run)


def clone_body(template, text: str):
    para = copy.deepcopy(template)
    for attr in list(para.attrib):
        if attr.endswith("paraId") or attr.endswith("textId"):
            del para.attrib[attr]
    for bookmark in list(para.findall(qn("w:bookmarkStart"))) + list(para.findall(qn("w:bookmarkEnd"))):
        para.remove(bookmark)
    set_paragraph(para, text)
    if text.startswith("资料依据"):
        ppr = para.get_or_add_pPr()
        for element in list(ppr.findall(qn("w:jc"))):
            ppr.remove(element)
        alignment = OxmlElement("w:jc")
        alignment.set(qn("w:val"), "left")
        ppr.append(alignment)
    return para


def fill_cell(cell, text: str) -> None:
    paragraphs = list(cell._tc.findall(qn("w:p")))
    reference = paragraphs[0]
    for paragraph in paragraphs:
        cell._tc.remove(paragraph)
    for line in text.split("\n"):
        cell._tc.append(clone_body(reference, line))


def after(anchor, body_template, lines: list[str]) -> None:
    for text in lines:
        paragraph = clone_body(body_template, text)
        anchor.addnext(paragraph)
        anchor = paragraph


def keep_with_next(paragraph):
    ppr = paragraph.get_or_add_pPr()
    if ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))


def write_package(source: Path, document, target: Path) -> dict:
    changed = {"word/document.xml": etree.tostring(document._element, xml_declaration=True, encoding="UTF-8", standalone=True)}
    with ZipFile(source) as original:
        # Original template author metadata is not part of the anonymous entry.
        for name in ("docProps/core.xml", "docProps/app.xml"):
            if name in original.namelist():
                xml = etree.fromstring(original.read(name))
                for node in xml.iter():
                    if etree.QName(node).localname in {"creator", "lastModifiedBy", "Company", "Manager"}:
                        node.text = ""
                changed[name] = etree.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)
        with ZipFile(target, "w", compression=ZIP_DEFLATED) as output:
            for entry in original.infolist():
                output.writestr(copy.copy(entry), changed.get(entry.filename, original.read(entry.filename)))
        with ZipFile(target) as output:
            preserved = [name for name in original.namelist() if name not in changed]
            assert all(original.read(name) == output.read(name) for name in preserved)
    return {"file": target.name, "sha256": digest(target), "template_sha256": digest(source),
            "changed_package_parts": sorted(changed), "preserved_package_parts": len(preserved)}


def make_form(source: Path, content: dict, metrics: dict | None):
    doc = Document(source)
    set_paragraph(doc.paragraphs[3]._p, "（申报正文草案）")
    table = doc.tables[0]
    replacements = {
        (0, 1): content["project_name"], (1, 1): "待补（匿名团队名称）",
        (3, 1): "待按平台口径填写", (3, 4): "待补", (4, 1): "待按平台口径填写",
        (6, 1): content["form"]["background"], (8, 1): content["form"]["idea"],
        (10, 1): content["form"]["solution"], (12, 1): content["form"]["benefits"],
        (13, 2): "创意设计类；基础研究类；\n软硬件开发类（选择）；工程实施类",
        (14, 2): "国际领先；国际先进；国内领先；国内先进；国内一般；无法判断（选择，尚无充分对标证据）",
        (15, 2): "□硬件 ☑软件 □工艺 ☑方法 □服务 □商业模式\n□其他",
        (16, 2): "□第1级 □第2级 □第3级 □第4级\n□第5级 □第6级 □第7级 □第8级 □第9级\n待依据成熟度定义与实施证据核定，暂不选择。",
        (17, 2): "同模型、同资料、同工具的固定流程（内部对照，非第三方产品）；结果见2026年10月3日小规模内部验收。其他国内外产品或技术的规范比较待补。",
        (19, 2): "新功能实现：可追溯核验、活动组合与双版包更新。\n质量、成本、效率及交付周期的提升值待实测，不填推断百分比。",
    }
    for (row, column), text in replacements.items():
        fill_cell(table.cell(row, column), text)
    for row in (5, 7, 9, 11):
        for paragraph in table.cell(row, 1).paragraphs:
            keep_with_next(paragraph._p)
    # Replace template writing-space height with content-driven rows, keeping
    # column structure, original type size, margins and merged sidebar intact.
    for row in (2, 6, 8, 10, 12):
        for height in list(table.rows[row]._tr.xpath("./w:trPr/w:trHeight")):
            height.getparent().remove(height)
    track = table.cell(2, 1)
    while len(track.paragraphs) > 1 and not track.paragraphs[-1].text.strip():
        element = track.paragraphs[-1]._p
        element.getparent().remove(element)
    for paragraph in track.paragraphs:
        if "智慧文旅与乡村振兴" in paragraph.text:
            set_paragraph(paragraph._p, paragraph.text.replace("5.定向赛道-智慧文旅与乡村振兴 □", "5.定向赛道-智慧文旅与乡村振兴 ☑"))
    nested = table.cell(18, 2).tables[0]
    rows = [
        ["完整交付", "待实测", "待实测", "不作推断", "同模型固定流程"],
        ["必要确认", "待实测", "待实测", "不作推断", "同模型固定流程"],
        ["正确拒绝", "待实测", "待实测", "不作推断", "同模型固定流程"],
        ["完整任务耗时", "待实测", "待实测", "不作推断", "同机固定流程"],
        ["交付约束违规", "待实测", "待实测", "不作推断", "逐项工具复核"],
    ]
    if metrics:
        rows = metrics["form_rows"]
        if len(rows) != 5 or any(len(row) != 5 for row in rows):
            raise ValueError("form_rows must be five rows by five columns")
    for row_i, row in enumerate(rows, 1):
        for col_i, text in enumerate(row):
            fill_cell(nested.cell(row_i, col_i), text)
    return write_package(source, doc, OUTPUT / "project-summary-draft.docx")


def make_description(source: Path, content: dict):
    doc = Document(source)
    paragraphs = list(doc.paragraphs)
    set_paragraph(paragraphs[3]._p, content["project_name"])
    body = copy.deepcopy(paragraphs[6]._p)
    for index in (5, 7, 8, 9, 10, 11, 13):
        keep_with_next(paragraphs[index]._p)
    # Blank template writing space includes a terminal page break, not a content section.
    for paragraph in paragraphs[14:]:
        paragraph._p.getparent().remove(paragraph._p)
    sections = [(6, "basis", True), (8, "overall", False), (9, "feasibility", False),
                (10, "innovation", False), (12, "implementation", True), (13, "prospects", False)]
    for index, key, replace in sections:
        anchor = paragraphs[index]._p
        lines = content["description"][key]
        if replace:
            set_paragraph(anchor, lines[0])
            after(anchor, body, lines[1:])
        else:
            after(anchor, body, lines)
    return write_package(source, doc, OUTPUT / "project-description-draft.docx")


def make_business(source: Path, content: dict):
    doc = Document(source)
    paragraphs = list(doc.paragraphs)
    set_paragraph(paragraphs[3]._p, content["project_name"])
    # The original subsection uses the same 14pt body face, 1.5 spacing, first-line indent.
    body = copy.deepcopy(paragraphs[8]._p)
    for index in range(5, 20):
        keep_with_next(paragraphs[index]._p)
    sections = [(5, "overview"), (6, "team"), (8, "features"), (9, "product_plan"),
                (11, "market"), (12, "competition"), (13, "risks"), (15, "development"),
                (16, "marketing"), (17, "revenue"), (18, "venture"), (19, "economics")]
    for index, key in sections:
        after(paragraphs[index]._p, body, content["business"][key])
    return write_package(source, doc, OUTPUT / "business-plan-draft.docx")


def check_lengths(content):
    counts = []
    for group, rules in LIMITS.items():
        for title, keys, limit in rules:
            blocks = []
            for key in keys:
                value = content[group][key]
                blocks.extend(value if isinstance(value, list) else [value])
            # Conservatively count Chinese, punctuation, digits and every Latin letter.
            count = len(re.sub(r"\s", "", "".join(blocks)))
            if count > limit:
                raise ValueError(f"{group}/{title}: {count}>{limit}")
            counts.append({"document": group, "section": title, "characters": count, "limit": limit})
    return counts


def write_readable(content, counts):
    labels = {"form": "项目简表", "description": "项目说明书", "business": "项目商业计划书"}
    lines = ["# 申报正文草案", "", "当前为按官方模板编排的正文草案，缺项及视觉验收状态见本目录README。", ""]
    for group, rules in LIMITS.items():
        lines.extend([f"## {labels[group]}", ""])
        for title, keys, limit in rules:
            count = next(c["characters"] for c in counts if c["document"] == group and c["section"] == title)
            lines.extend([f"### {title}", "", f"正文计数：{count}/{limit}（含标点及逐字符计入英文，较通常字数口径保守）。", ""])
            for key in keys:
                values = content[group][key]
                for value in values if isinstance(values, list) else [values]:
                    lines.extend([value, ""])
    (OUTPUT / "submission-body.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--templates-dir", type=Path, default=Path.home() / "Downloads")
    parser.add_argument("--metrics", type=Path, help="Explicit reviewed measurements; never inferred by builder")
    args = parser.parse_args()
    content = json.loads((OUTPUT / "draft-content.json").read_text(encoding="utf-8"))
    metrics = json.loads(args.metrics.read_text(encoding="utf-8")) if args.metrics else None
    sources = {}
    for index in range(4):
        pattern = "附件：*智慧城市*指南.pdf" if index == 0 else f"附件{index}*智慧城市*.docx"
        matches = list(args.templates_dir.glob(pattern))
        if len(matches) != 1:
            raise ValueError(f"Expected exactly one official input for {pattern}")
        source = matches[0]
        if digest(source) != EXPECTED_HASHES[index]:
            raise ValueError(f"Official template hash changed: attachment {index}")
        sources[index] = source
    counts = check_lengths(content)
    documents = [make_form(sources[1], content, metrics), make_description(sources[2], content),
                 make_business(sources[3], content)]
    write_readable(content, counts)
    manifest = {"status": "draft-awaiting-external-evidence-and-canonical-render", "date": content["version_date"],
                "source_hashes": EXPECTED_HASHES, "outputs": documents, "character_counts": counts,
                "count_method": "non-whitespace Unicode characters, including punctuation and individual Latin letters",
                "metrics_source": args.metrics.name if args.metrics else None,
                "originals_unchanged": all(digest(sources[i]) == EXPECTED_HASHES[i] for i in range(4))}
    (OUTPUT / "build-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"outputs": documents, "counts": counts}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
